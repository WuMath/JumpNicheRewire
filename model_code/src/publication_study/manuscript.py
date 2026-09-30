"""Build a factual Word manuscript ONLY after every planned experiment has results.

Call with the Codex bundled Python after the document skill's authoring marker.
The resulting manuscript still requires figure/source review and render QA.
"""
from pathlib import Path
import csv
import json
from statistics import mean
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
RESULT=ROOT/'data/result/publication_20260920'
PROCESSED=ROOT/'data/processed/publication_20260920'

def read_csv(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as stream:return list(csv.DictReader(stream))

def number(value,digits=4):
    try:return f'{float(value):.{digits}g}'
    except (ValueError,TypeError):return '未定义'

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    if widths is None:
        widths={2:[1.65,5.15],3:[1.65,2.55,2.60],4:[2.0,1.6,1.6,1.6],5:[1.4,1.25,1.05,2.15,.95]}[len(headers)]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
    props.append(borders)
    for cell,value in zip(t.rows[0].cells,headers):cell.text=str(value)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,values):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=1
            cp=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:w'),'80');item.set(qn('w:type'),'dxa');margins.append(item)
            cp.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');cp.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.1
                if ri==0:p.paragraph_format.keep_with_next=True
                if len(cell.text)<15:p.alignment=1
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
    doc.add_paragraph()

CASE_DISCUSSION={
 'case_mouse_brain':[
 '空间筛选后仅Lmo4与Egr1保留在调控者目录，脑网络因此是有限候选面板，不是全脑TF网络。排名第一的展示链为Cntn1–Nrcam–Lmo4–Penk，最高path_score为0.01295，delta_beta为−0.004391。这里的箭头是算法关联分解，不证明Nrcam直接调控Lmo4，也不证明该链在该spot产生了因果效应。',
 '全部通路检验q值均为1。Dravet、Rett和心肌钙调控等数据库条目名称不能用于诊断组织疾病或宣称通路激活。当前结果的实际用途是提供带空间位置的Penk等候选以供后续原位表达核对；没有通过统计阈值的通路结论。脑区未获得独立组织学标注，不能把expression_cluster编号解释为特定脑区。'],
 'case_mouse_kidney':[
 '102项检验中3项显著，均为同一条Adipogenesis Genes WP447，分别来自global_GRN、TF_Hnf4a与TF_Irx1候选集合，不是3条独立通路。每组前50靶基因中命中8个基因，背景147基因中该通路也仅有8个，p=0.00012051，BH q=0.00409735。重叠为Adipoq、Cfd、Fabp4、Pck1、Plin1、Retn、Scd1与Ucp1。',
 '该组合提示脂肪相关表达成分，但不能推断肾上皮发生脂肪分化。组织中不同细胞成分、周围组织混入以及有限高变基因面板均是待检查的解释；目前没有细胞组成反卷积或组织学证据判定来源。PPAR条目的最佳q约0.0822，未达阈值。',
 '最高展示链为Calr–Itgav–Hnf4a–Mep1b，path_score=0.01647，delta_beta=−0.005451。Chen等通过遗传学与ChIP-seq显示HNF4参与肾刷状缘转录程序[15]，支持Hnf4a作为组织相关候选的合理性，但并未验证本研究这条Calr起始链。后续可先定位Hnf4a与Mep1b表达，再用独立受体阻断和TF扰动检验模型方向；这些实验尚未实施。'],
 'case_human_breast_cancer':[
 '仅AEBP1与FOS进入最终调控面板。420项通路检验无显著结果；最佳q=0.16216。LR_CALR__ITGAV的两个候选COL1A1、COL1A2命中黏附和PI3K相关条目，但样本集合仅含2基因且未通过BH校正，不应据此宣称这些通路受到激活。数据库出现胰腺癌或高血糖条目也不表示切片具有相应疾病。',
 '最高展示链为SPP1–ITGB1–AEBP1–CTSL，path_score=0.01925，delta_beta=−0.008306。它可用作空间共定位及受体复合物核对的候选；单独的ITGB1不能代表完整整合素异二聚体。本次没有蛋白水平测量证实复合物。已有乳腺癌单细胞和空间图谱显示微环境细胞组成具有空间异质性[16]；这说明需要细胞组成校正，而非为当前链提供独立验证。']}

def build():
    audit=json.loads((RESULT/'scientific_audit.json').read_text(encoding='utf-8'))
    if not audit.get('complete_experiments'):
        raise RuntimeError('Cannot produce a complete experimental paper while scheduled experiments are missing or failed')
    metrics=read_csv(RESULT/'03_evaluation/dataset_metrics.csv')
    cases=read_csv(RESULT/'04_cases/case_summary.csv')
    registry=read_csv(PROCESSED/'benchmark_registry.csv')
    claim=json.loads((RESULT/'03_evaluation/claim_audit.json').read_text(encoding='utf-8'))
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.85)
    for style in ('Normal','Title','Heading 1','Heading 2','Caption'):
        st=doc.styles[style];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    # python-docx's bundled default template can contain a blue Title border.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.25
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(21)
    for h in ('Heading 1','Heading 2'):doc.styles[h].paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0];footer.alignment=2
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.add_paragraph('JumpNicheRewire空间调控网络推断的公开数据评估与组织通路分析','Title')
    doc.add_paragraph('公开数据核心范围评估及三项组织切片应用')
    doc.add_heading('摘要',1)
    superior=claim.get('all_metrics_strictly_better')
    conclusion=('在预定的有噪声测试集合上，所检查指标的条件组平均值均高于两个对照。统计显著性和区间需分别见结果，均值领先不等于每个数据集都领先。'
        if superior else '本次实验未支持全部评价指标均优于两个对照的结论；结果保留各方法的优势与不足，不据此宣称全面领先。')
    doc.add_paragraph('空间转录组为研究局部转录调控和细胞间通信提供了表达及位置数据。我们将JumpNicheRewire的空间图、稀疏调控回归、软先验和通信条件重构连接为可审计实现，并使用scMultiSim作者已公开发布的数据开展独立真值评价。预定范围包括6个运行验收数据集、18个有测量噪声的留出数据集及2个较大空间模拟对象；后两者为公开的潜在真实计数，单独报告。比较方法为官方Arboreto实现的GENIE3和GRNBoost2，主要共同任务为TF到非TF靶基因的全局网络恢复。'+conclusion)
    doc.add_paragraph('我们进一步分析小鼠脑、小鼠肾和人乳腺肿瘤三个公开Visium FFPE切片，保留各切片全部组织内spot，并将推断网络投射到物种匹配的WikiPathways集合。通路分析采用实际拟合靶基因背景、超几何检验和BH校正。案例结果为先验参与的探索性关联，不代表通路激活的独立实验验证，也不支持干预因果结论。')
    doc.add_paragraph('18份有噪声留出数据的平均AP分别为JumpNicheRewire 0.2351、GENIE3 0.2220、GRNBoost2 0.2577。多重校正后的配对检验均未达到0.05。真实案例仅小鼠肾的一个通路在三个候选集合中通过FDR阈值；脑和乳腺案例无显著富集。三个案例最终仅拟合2、3、2个TF，限制了网络的覆盖范围。')
    doc.add_paragraph('关键词  空间转录组  基因调控网络  配体受体  图融合  公开数据评估  通路富集')
    doc.add_heading('引言',1)
    doc.add_paragraph('基因调控网络推断需要区分表达相关性、条件依赖和实验因果关系。GENIE3通过逐靶基因树模型估计候选调控者的重要性[1]；GRNBoost2采用梯度提升和早停机制，并由Arboreto提供可扩展执行框架[2]。它们是有效的表达型GRN对照，但不直接输出细胞空间网络或完整的配体到靶基因通路。')
    doc.add_paragraph('本研究考察引入空间邻近和通信条件项能否改善共同GRN任务，并展示其在组织切片中的机制候选分析用途。我们不把对照不支持的输出算作零分，也不把公开模拟标签与真实组织中的通路知识混为一谈。BEELINE说明了统一任务和评价背景对GRN算法比较的重要性[3]。')
    doc.add_heading('材料与方法',1)
    doc.add_heading('数据来源和纳入范围',2)
    table(doc,['数据','原始范围','本研究用途'],[
        ['scMultiSim main_tree1/3/5','500细胞、110基因、2个噪声设置、4个重复','6份验收和18份留出评价'],
        ['scMultiSim spatial 5000','5000细胞、500基因','较大空间对象的独立描述性评价'],
        ['scMultiSim spatial 3000','3000细胞、2000基因','较大空间对象的独立描述性评价'],
        ['10x成年小鼠脑FFPE','2264个组织spot','案例一'],['10x成年小鼠肾FFPE','3124个组织spot','案例二'],
        ['10x人乳腺肿瘤FFPE','2518个组织spot','案例三']])
    doc.add_paragraph('模拟数据来自Zenodo 13119261，关联scMultiSim原始研究[4,5]。压缩包MD5与发布页一致，并记录文件内容哈希。本轮不调用模拟器生成新的正式数据。我们并未声称纳入Zenodo的全部数据；所选范围在模型评价前记录。所有细胞均保留，大面板最多按表达方差选择200个非TF靶基因，共同输入及评价范围对三个方法相同。详细原始和保留维度见数据注册表。')
    doc.add_paragraph('三个真实数据分别取自10x官方发布页[6–8]，脑切片复用本地已下载文件，另外两例重新下载并记录SHA256。它们来自不同组织，但每个案例仅含一个切片，不能将多个spot视为独立生物学样本。每例预处理保留600基因、12个表达TF及最多40个严格LR候选，模型空间筛选及回归目标上限为150；实际拟合数见案例结果。')
    doc.add_paragraph('空间筛选还会将TF目录与保留的空间基因集合取交集（spatial.py的candidate_tfs分支），所以预处理的12个TF不等于最终拟合12个TF。本轮保持冻结实现，未依据测试结果扩大面板。大空间对象除200个评价target外还保留LR身份基因，因此拟合target数可能超过200；这些额外基因不改变正式评价候选全集。')
    doc.add_heading('模型实现和证据边界',2)
    doc.add_paragraph('模块一对标准化表达计算空间统计，构建基于空间距离与表达状态的邻接图。置换次数固定为999。模块二使用稀疏多变量回归和图融合惩罚估计细胞分辨率的基础TF–gene系数，采用ADMM求解。模块三将TF–gene与LR–TF参考证据转换为软先验。模块四由邻近发送spot的配体和接收spot的受体表达构建暴露，并使用分区共享的交互系数估计通信条件重构，FISTA处理相应惩罚项。外层联合迭代更新通信状态，使用未阻尼固定点残差而非阻尼后变化判断收敛。')
    doc.add_paragraph('实现的系数分解为：最终TF–gene系数等于基础系数与各LR暴露乘其调节系数之和。该分解提供关联解释；没有独立干预时，不将其等同于已识别的因果效应。原始V5文档包含更广泛的方法设计，本研究只报告当前代码实际执行的模块，不声称所有拟议增强层均已验证。')
    doc.add_paragraph('本轮固定L1惩罚0.02、空间融合惩罚0.15、重构group惩罚0.002和ridge惩罚0.0001，随机种子20260920。表达先按每细胞总量归一化至10000后log1p变换；回归内部的中心化与尺度变换见utils.py和base_grn.py。两个树模型均采用单工作进程官方实现。每项的完整参数及求解容差保存于model_config.json，冻结文件保存设置与数值实现哈希。')
    table(doc,['模块','主要实现文件'],[['空间图与筛选','src/jump_niche_rewire/spatial.py'],['基础GRN与ADMM','base_grn.py  solvers.py'],['软先验','prior.py'],['通信与重构','communication.py  rewiring.py'],['整体调用与收敛','pipeline.py  convergence.py'],['本研究运行与评价','src/publication_study/workflow.py  analysis.py']])
    doc.add_paragraph('模拟benchmark不提供GRN边真值作为先验；给模型的TF身份来自作者候选调控者目录，LR身份来自作者声明的候选参数，不按活跃通信标签筛选。真实案例使用既有物种匹配参考；人类LR经过胞外配体与跨膜受体角色、方向及文献来源检查，复杂体不能可靠映射时排除。人类案例无匹配motif证据时明确缺失。LR→TF路径来自OmniPath有向信号参考，只表示参考路径支持。')
    doc.add_heading('共同任务和参数隔离',2)
    doc.add_paragraph('共同评价为TF→非TF target的有向无符号GRN。JumpNicheRewire在完整数值系数上对细胞维度取绝对值最大值，输出global_grn.parquet；评价不读取已经截断的展示网络。GENIE3与GRNBoost2使用官方Arboreto重要性分数。未输出的候选边赋0，所有方法使用相同候选全集。仅大小写匹配不被称为同源转换。')
    doc.add_paragraph('使用既有参数作为预先规定的设置，而非利用test选择参数。6份validation用于检查运行与收敛，全部完成后保存参数与实现哈希，再运行20份test。对每次任务分别记录输入哈希、软件版本、耗时、峰值RSS、异常和收敛率。至少90%的适用target求解应收敛，且适用外层联合迭代必须通过。失败数据不作为有效预测进入指标。')
    doc.add_heading('评价指标与统计',2)
    table(doc,['指标','定义与意义','依据'],[
        ['AP','按召回增量加权精确率，关注真边整体排序','[3,9,10]'],
        ['AUROC','随机正边高于随机负边的概率，同分计0.5','[9,11]'],
        ['Precision@k','前k候选中的真边比例，反映验证效率','[3,10]'],
        ['Recall@k','前k候选找回的真边比例，反映遗漏','[10]'],
        ['F1@k','精确率与召回率的调和平均','[10]'],
        ['Early precision','k等于正标签数时的精确率，不是EPR比值','[3]']])
    doc.add_paragraph('k预先固定为10、50、100。同分整体进入AP阈值；top-k切入同分组时，按该组均匀排列的期望真阳性数计算，避免输入行顺序影响分数[12]。AP不同于梯形PR面积；正例占比是数据诊断，不是模型可争取领先的性能指标。')
    doc.add_paragraph('有噪声的18份测试数据按tree×sigma组成6个条件组，在组内先平均重复间的配对差值，然后对条件组bootstrap 2000次得到百分位区间，同时进行双侧Wilcoxon符号秩检验和Holm多重比较校正。由于模拟共享网络结构，区间主要反映该预定生成范围的稳定性，不能代表对所有GRN拓扑的普遍性能。较大潜在真实计数对象单独展示，不与有测量噪声结果混合掩盖差异。')
    doc.add_heading('真实案例与通路统计',2)
    doc.add_paragraph('使用PCA10与KMeans6产生表达计算分区，明确其不是细胞类型或经组织学验证的脑区。以实际拟合的靶基因为富集背景，按全局GRN、单TF及LR机制候选分数选择最多50个正分靶基因。使用物种专属WikiPathways 2024集合[13,14]，计算超几何上尾概率，包含零重叠检验，对每个案例的全部候选集合与通路检验统一作BH校正。q≤0.05标记显著。先验参与的富集仅解释已推断网络，不作为独立因果验证。')
    doc.add_heading('结果',1)
    doc.add_heading('运行完整性和数值收敛',2)
    states=[json.loads(p.read_text(encoding='utf-8')) for p in (RESULT/'02_runs').rglob('run_status.json') if 'attempt_history' not in p.parts]
    accepted=[s for s in states if s.get('state')=='completed']
    doc.add_paragraph(f'本次清单包含26份模拟数据上的78次方法运行及3次真实案例运行。已完成任务数为{len(accepted)}。测试集正式接受{claim["accepted_test_runs"]}/{claim["expected_test_runs"]}次运行。运行失败、收敛和资源记录见run_audit.csv与各任务run_status.json。')
    bench_states=[s for s in accepted if s['split']!='case']
    table(doc,['方法','模拟完成数','平均耗时秒','已知最大RSS GiB'],[[m,sum(s['method']==m for s in bench_states),
        number(mean([s['runtime_seconds'] for s in bench_states if s['method']==m])),
        number(max([s['peak_rss_gib'] for s in bench_states if s['method']==m and s.get('peak_rss_gib') is not None]))] for m in ('JumpNicheRewire','GENIE3','GRNBoost2')])
    doc.add_paragraph('上表只比较各方法相同的26份模拟任务，不把三个真实案例的耗时混入JumpNicheRewire平均值。一次GENIE3孤儿进程完成后经输入、实现和输出哈希核验复用，但该次完整峰值RSS缺失；已知最大值不等于所有运行完整峰值。资源监测文件重命名竞争及Windows编码曾导致工程失败，失败记录保留在attempt_history，未重复计入81项。')
    doc.add_heading('公共模拟GRN恢复',2)
    test=[r for r in metrics if r['split']=='test' and r['stratum']!='large_spatial_latent_counts']
    names=['AP','AUROC','Early_precision','Precision_at_10','Recall_at_10','F1_at_10','Precision_at_50','Recall_at_50','F1_at_50','Precision_at_100','Recall_at_100','F1_at_100']
    methods=('JumpNicheRewire','GENIE3','GRNBoost2')
    table(doc,['指标',*methods],[[name,*[number(mean([float(r['value']) for r in test if r['metric']==name and r['method']==m])) for m in methods]] for name in names])
    doc.add_paragraph(conclusion+'完整逐数据集结果和每项指标独立表格保留在03_evaluation/per_metric；条件组配对差值与置信区间见paired_statistics.csv。不能根据平均名次隐藏任何预定指标。')
    figure=RESULT/'03_evaluation/figures/AP.png'
    if figure.exists():doc.add_picture(str(figure),width=Inches(5.8));doc.add_paragraph('图1  留出数据的AP分布与逐数据集配对关系。黑色短线表示均值。','Caption')
    paired=read_csv(RESULT/'03_evaluation/paired_statistics.csv')
    table(doc,['指标','对照','平均差值','95%区间','Holm p'],[[r['metric'],r['comparator'],number(r['mean_paired_difference']),
        number(r['ci_low'])+' 至 '+number(r['ci_high']),number(r.get('holm_p'))] for r in paired if r['metric'] in ('AP','AUROC','F1_at_100')])
    doc.add_paragraph('全部24项配对比较的Holm p均不小于0.75。部分未校正bootstrap区间不跨0，但它们与多重校正Wilcoxon检验回答的问题不同，不能据此越过多重检验门槛。只有6个条件组、共享GRN结构和有限重复使区间与检验效能受限。')
    doc.add_heading('较大空间模拟对象',2)
    large=[r for r in metrics if r['split']=='test' and r['stratum']=='large_spatial_latent_counts' and r['metric'] in ('AP','AUROC','F1_at_100')]
    table(doc,['数据','方法','指标','结果'],[[r['dataset_id'].replace('published_spatial_',''),r['method'],r['metric'],number(r['value'])] for r in large])
    doc.add_paragraph('这些对象保留全部3000或5000细胞，但靶基因范围受预先设定的表达面板限制。它们缺少本轮重新生成的技术噪声，因此不能据其数值直接推断真实组织中的网络准确率。')
    labels={'case_mouse_brain':'小鼠脑','case_mouse_kidney':'小鼠肾','case_human_breast_cancer':'人乳腺肿瘤'}
    for i,row in enumerate(cases,1):
        key=row['dataset_id'];folder=RESULT/'04_cases'/key
        doc.add_heading('案例'+str(i)+' '+labels[key],2)
        doc.add_paragraph(f'本例包含{row["spots"]}个组织spot，实际拟合{row["fitted_tfs"]}个TF和{row["fitted_targets"]}个靶基因。完成{row["pathway_tests"]}项通路检验，其中{row["significant_tests"]}项通过预定BH阈值。所有显著与不显著检验均保留，避免只挑选符合组织先验的通路。')
        terms=read_csv(folder/'pathway_enrichment_all_tests.csv')
        table(doc,['候选集合','通路','重叠','BH q'],[[r['gene_set'],r['pathway'],r['overlap_genes'],number(r['q_value'])] for r in terms[:3]],widths=[1.35,2.65,2.1,.7])
        for paragraph in CASE_DISCUSSION[key]:doc.add_paragraph(paragraph)
        fig=folder/'case_overview.png'
        if fig.exists():doc.add_picture(str(fig),width=Inches(5.6));doc.add_paragraph(f'图{i+1}  {labels[key]}推断调节子活性的空间分布及全局GRN通路富集。红虚线为q=0.05；q=1对应长度为0的条形。空间图展示模型推断，不是蛋白活性测量。','Caption')
        if int(row['significant_tests'])==0:
            doc.add_paragraph('本例未检出通过预定阈值的富集项，不能将排名最前的通路写成已获得显著支持。应结合更大靶基因面板、独立样本和测量证据进一步检查。')
        else:
            doc.add_paragraph('富集结果说明模型候选在给定靶基因背景下对相应功能集合有所集中。它不证明这些通路被激活，也不证明通信引起了该变化。mechanism_paths_preview.csv提供LR–TF–target候选链及其数值贡献，便于优先设计受体阻断、TF干预或空间蛋白验证。')
    doc.add_heading('讨论',1)
    doc.add_paragraph(conclusion+'空间正则项可能改善局部估计，也可能在不适当的邻域或混杂结构下引入偏差。全局树模型擅长某些非线性表达关系，因此引入空间信息本身并不保证超过表达型方法。单次强结果不足以证明普遍有效。')
    doc.add_paragraph('本轮比较范围仅包括共同GRN任务和两个表达型对照，未覆盖所有空间方法，也未完成LR–TF–gene整条机制的独立真值恢复。要支持更强的模型方法学结论，仍需空间对照、模块消融、多网络拓扑和独立扰动评价。不能将本研究核心范围报告冒充这些扩展实验已经完成。')
    doc.add_paragraph('三个案例展示从空间表达、调控网络到通路候选的可重复分析流程。对生物学家的直接用途是给出可追溯的受体、TF和靶基因优先级及对应组织位置，减少完全无依据的候选筛选。但每例单切片、计算分区、有限基因面板和先验依赖限制了结论的外推。研究需要通过独立样本、细胞组成测量和干预验证继续检验这些候选。')
    doc.add_heading('数据和代码可用性',1)
    doc.add_paragraph('所有正式数据源链接、压缩包校验码和下载文件哈希保存在01_data。原始数据不随模型代码重复打包。入口为main_jumpniche.py和preprocess_jumpniche.py，配置类位于configs/public_sim_data/publication_study.py。源码、配置、运行协议与测试可由本地发布包复现；硬件差异可能影响耗时，不能保证任意平台逐位相同。')
    doc.add_heading('参考文献',1)
    refs=[
      'Huynh-Thu VA et al. Inferring Regulatory Networks from Expression Data Using Tree-Based Methods. PLoS ONE, 2010. https://doi.org/10.1371/journal.pone.0012776',
      'Moerman T et al. GRNBoost2 and Arboreto: efficient and scalable inference of gene regulatory networks. Bioinformatics, 2019. https://doi.org/10.1093/bioinformatics/bty916',
      'Pratapa A et al. Benchmarking algorithms for gene regulatory network inference from single-cell transcriptomic data. Nature Methods, 2020. https://doi.org/10.1038/s41592-019-0690-6',
      'Li H et al. scMultiSim: simulation of single-cell multi-omics and spatial data guided by gene regulatory networks and cell–cell interactions. Nature Methods, 2025. https://doi.org/10.1038/s41592-025-02651-0',
      'Li H. scMultiSim Benchmarking Datasets. Zenodo, 2024. https://doi.org/10.5281/zenodo.13119261',
      '10x Genomics. Adult Mouse Brain FFPE. https://www.10xgenomics.com/datasets/adult-mouse-brain-ffpe-1-standard-1-3-0',
      '10x Genomics. Adult Mouse Kidney FFPE. https://www.10xgenomics.com/datasets/adult-mouse-kidney-ffpe-1-standard-1-3-0',
      '10x Genomics. Human Breast Cancer Ductal Carcinoma In Situ Invasive Carcinoma FFPE. https://www.10xgenomics.com/datasets/human-breast-cancer-ductal-carcinoma-in-situ-invasive-carcinoma-ffpe-1-standard-1-3-0',
      'Saito T, Rehmsmeier M. The Precision-Recall Plot Is More Informative than the ROC Plot When Evaluating Binary Classifiers on Imbalanced Datasets. PLoS ONE, 2015. https://doi.org/10.1371/journal.pone.0118432',
      'Manning CD, Raghavan P, Schütze H. Introduction to Information Retrieval. Cambridge University Press, 2008. Chapter 8. https://nlp.stanford.edu/IR-book/html/htmledition/evaluation-of-ranked-retrieval-results-1.html',
      'Fawcett T. An introduction to ROC analysis. Pattern Recognition Letters, 2006. https://doi.org/10.1016/j.patrec.2005.10.010',
      'McSherry F, Najork M. Computing Information Retrieval Performance Measures Efficiently in the Presence of Tied Scores. ECIR, 2008. https://www.microsoft.com/en-us/research/publication/computing-information-retrieval-performance-measures-efficiently-in-the-presence-of-tied-scores/',
      'Agrawal A et al. WikiPathways 2024: next generation pathway database. Nucleic Acids Research, 2024. https://doi.org/10.1093/nar/gkad960',
      'Kuleshov MV et al. Enrichr: a comprehensive gene set enrichment analysis web server 2016 update. Nucleic Acids Research, 2016. https://doi.org/10.1093/nar/gkw377',
      'Chen L et al. The nuclear receptor HNF4 drives a brush border gene program conserved across murine intestine, kidney, and embryonic yolk sac. Nature Communications, 2021, 12:2886. https://doi.org/10.1038/s41467-021-22761-5',
      'Wu SZ et al. A single-cell and spatially resolved atlas of human breast cancers. Nature Genetics, 2021, 53:1334–1347. https://doi.org/10.1038/s41588-021-00911-1',
    ]
    for i,ref in enumerate(refs,1):doc.add_paragraph(f'[{i}] {ref}')
    output=RESULT/'05_manuscript/JumpNicheRewire_public_data_study.docx'
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    return output

if __name__=='__main__':print(build())
