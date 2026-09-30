"""评价结果的逐指标导出。不读取输入表达/真值，也不调用模型。"""
from __future__ import annotations
from pathlib import Path
import numpy as np
import pandas as pd
from .io import safe_name
from .metric_references import definition,bibliography

def _workbook(path,sheets):
    # 项目运行时的自动导出，沿用现有 pandas/openpyxl 环境。
    # 内容全是已算好的科学指标，不把中间统计冒充可编辑 Excel 公式。
    from openpyxl.styles import Font,PatternFill,Alignment
    with pd.ExcelWriter(path,engine="openpyxl") as writer:
        for name,frame in sheets.items():
            frame.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]
            ws.freeze_panes="B2" if name=="by_dataset" else "A2"
            if ws.max_column and ws.max_row>1:
                ws.auto_filter.ref=ws.dimensions
            ws.row_dimensions[1].height=30
            for cell in ws[1]:
                cell.font=Font(bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor="26465E")
                cell.alignment=Alignment(vertical="center",wrap_text=True)
            for col_index,column in enumerate(frame.columns,1):
                letter=ws.cell(1,col_index).column_letter
                text_size=max([len(str(column))]+[len(str(v)) for v in frame[column].head(100) if pd.notna(v)])
                ws.column_dimensions[letter].width=min(75,max(18,text_size+2))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(vertical="top",wrap_text=name in {"definition","references","audit"})
                    if isinstance(cell.value,float):
                        cell.number_format="0.000000"
                if name in {"definition","references"}:
                    ws.row_dimensions[row[0].row].height=78

def export_metric_tables(frame,audit,selected_rows,ec,output):
    """任务×指标独立成表；缺失方法和数据保留，不填伪造的零。"""
    from .comparison import mean_interval
    methods=list(ec.methods)
    ids=[r["dataset_id"] for r in selected_rows]
    if len(methods)!=len(set(methods)) or len(ids)!=len(set(ids)):
        raise ValueError("所选方法/数据集编号必须唯一")
    manifest=[]
    if frame.empty:
        return pd.DataFrame()
    for (task,metric),group in frame.groupby(["task","metric"],sort=True):
        if group.duplicated(["dataset_id","method"]).any():
            raise ValueError(f"重复指标记录：{task}/{metric}")
        wide=group.pivot(index="dataset_id",columns="method",values="value").reindex(index=ids,columns=methods)
        wide.columns.name=None
        long=wide.reset_index().melt(id_vars="dataset_id",var_name="method",value_name=metric)
        detail=group.drop(columns=["value","task","metric"])
        long=long.merge(detail,on=["dataset_id","method"],how="left",validate="one_to_one")
        audit_view=long[["dataset_id","method",metric]].merge(audit,on=["dataset_id","method"],how="left",validate="one_to_one")
        audit_view["metric_available"]=audit_view[metric].notna()
        audit_view["metric_note"]=np.where(audit_view.metric_available,"available",
            np.where(audit_view.state.eq("completed"),"task unsupported or metric undefined; not zero","see run state; not zero"))
        summaries=[]
        matched=[]
        common=wide.dropna(axis=0,how="any")
        for method in methods:
            values=wide[method].dropna().to_numpy(dtype=float)
            low,high=mean_interval(values,ec,np.random.default_rng(ec.random_seed))
            summaries.append(dict(method=method,**{metric:float(values.mean()) if len(values) else np.nan},
                                  ci_lower=low,ci_upper=high,evaluated_datasets=len(values),
                                  selected_datasets=len(ids),missing_or_undefined=len(ids)-len(values),
                                  averaging="macro mean over available datasets; subsets may differ"))
            shared=common[method].to_numpy(dtype=float)
            low,high=mean_interval(shared,ec,np.random.default_rng(ec.random_seed))
            matched.append(dict(method=method,**{metric:float(shared.mean()) if len(shared) else np.nan},
                                ci_lower=low,ci_upper=high,common_datasets=len(shared),
                                averaging="same complete datasets for every selected method"))
        summary=pd.DataFrame(summaries)
        matched=pd.DataFrame(matched)
        info=definition(metric,ec)
        # 竖向定义表避免一行过宽；保留原始 URL 便于复制到论文文献管理软件。
        information=dict(task=task,**info,confidence_level=ec.confidence_level,
                         repetition_unit="dataset/seed, not edges; SERGIO shared topology is conditional replication",
                         missing_policy="blank/NA, never invented zero; see audit",
                         summary_policy="method_summary may use different subsets; matched_summary uses all-method intersection")
        definition_frame=pd.DataFrame(dict(field=list(information),value=list(information.values())))
        cited=info["reference_ids"].split(", ")+["EFRON1979","WILCOXON1945","HOLM1979"]
        sheets=dict(by_dataset=wide.reset_index(),method_summary=summary,matched_summary=matched,
                    long=long,audit=audit_view,definition=definition_frame,references=bibliography(cited))
        directory=Path(output)/"by_metric"/safe_name(task)
        directory.mkdir(parents=True,exist_ok=True)
        stem=safe_name(metric)
        workbook=directory/f"{stem}.xlsx"
        _workbook(workbook,sheets)
        for suffix,part in (("",sheets["by_dataset"]),("_summary",summary),("_matched",matched),("_long",long)):
            part.to_csv(directory/f"{stem}{suffix}.csv",index=False,encoding="utf-8-sig")
        manifest.append(dict(task=task,metric=metric,workbook=str(workbook),wide_csv=str(directory/f"{stem}.csv"),
                             summary_csv=str(directory/f"{stem}_summary.csv"),common_datasets=len(common),
                             reference_ids=info["reference_ids"]))
    index=pd.DataFrame(manifest)
    index.to_csv(Path(output)/"metric_table_manifest.csv",index=False,encoding="utf-8-sig")
    return index
